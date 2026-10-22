# -*- coding: utf-8 -*-
"""
check_ddl.py — DDL(ddl.sql) 정적 검사기

MySQL 에 실행하기 전에, 파일만 읽고 세 가지를 대조한다.
  ① 4일차 TO-BE 모델 — 테이블 4개, PK, FK 가 맞는가
  ② 2일차 표준 사전 — 컬럼 이름이 「용어」 물리명이고, 자료형·길이가 「도메인」 과 같은가
  ③ 데이터 — 실제로 비어 있는 값이 들어올 컬럼이 NULL 을 허용하는가

실행 (C:\\dstd 에서)
  .\\.venv\\Scripts\\python.exe src\\check_ddl.py                     docs\\ddl.sql
  .\\.venv\\Scripts\\python.exe src\\check_ddl.py templates\\ddl_템플릿.sql
"""
from __future__ import annotations

import argparse
import re
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
from check_dictionary import ROOT, Report   # noqa: E402

TABLES = {"CUST": ["CUST_NO"], "PRD": ["PRD_NO"], "ORD": ["ORD_NO"], "ORD_ITEM": ["ORD_NO", "ORD_SEQ"]}
FKS = [("ORD", "CUST_NO", "CUST"), ("ORD_ITEM", "ORD_NO", "ORD"), ("ORD_ITEM", "PRD_NO", "PRD")]
# 데이터에 실제로 빈 값이 들어오는 컬럼 — NOT NULL 이면 적재가 실패한다
MUST_NULL = {
    ("ORD", "CUST_NO"): "비회원 주문 (원천 76행, 업무설명서 2.1)",
    ("CUST", "EMAIL_ADDR"): "이메일 없는 고객 36명",
    ("ORD", "CNCL_DT"): "취소되지 않은 주문",
    ("PRD", "TAX_TP_CD"): "정의 없는 과세구분 Z 상품 2개",
    ("ORD", "PAY_MTHD_CD"): "결제수단 ETC (확인필요)",
}
TYPE_ALIAS = {"INTEGER": "INT"}

CREATE_RE = re.compile(r"CREATE\s+TABLE\s+(?:IF\s+NOT\s+EXISTS\s+)?`?(\w+)`?\s*\(", re.I)
COL_RE = re.compile(r"^`?(\w+)`?\s+(\w+)\s*(?:\(\s*(\d+)\s*(?:,\s*(\d+))?\s*\))?(.*)$", re.I | re.S)
PK_RE = re.compile(r"PRIMARY\s+KEY\s*\(([^)]*)\)", re.I)
FK_RE = re.compile(r"FOREIGN\s+KEY\s*\(([^)]*)\)\s*REFERENCES\s+`?(\w+)`?\s*\(([^)]*)\)", re.I)


def strip_comments(sql: str) -> str:
    sql = re.sub(r"/\*.*?\*/", " ", sql, flags=re.S)
    return "\n".join(re.sub(r"(--|#).*$", "", ln) for ln in sql.splitlines())


def split_defs(body: str):
    """괄호 안의 콤마는 무시하고 정의를 나눈다. 따옴표 안도 무시."""
    out, cur, depth, q = [], "", 0, None
    for ch in body:
        if q:
            cur += ch
            if ch == q:
                q = None
            continue
        if ch in "'\"":
            q = ch
        elif ch == "(":
            depth += 1
        elif ch == ")":
            depth -= 1
        if ch == "," and depth == 0:
            out.append(cur.strip())
            cur = ""
        else:
            cur += ch
    if cur.strip():
        out.append(cur.strip())
    return out


def bodies(sql: str):
    """CREATE TABLE 이름 ( ... ) 의 괄호 안을 짝을 맞춰 잘라 낸다 (따옴표 안 괄호는 무시)."""
    for m in CREATE_RE.finditer(sql):
        i, depth, q = m.end(), 1, None
        while i < len(sql) and depth:
            ch = sql[i]
            if q:
                q = None if ch == q else q
            elif ch in "'\"":
                q = ch
            elif ch == "(":
                depth += 1
            elif ch == ")":
                depth -= 1
            i += 1
        yield m.group(1).upper(), sql[m.end():i - 1]


def parse(sql: str, rep: Report):
    tables = {}
    for name, body in bodies(sql):
        t = {"cols": {}, "pk": [], "fks": []}
        for d in split_defs(body):
            u = d.upper()
            if u.startswith("PRIMARY KEY"):
                t["pk"] = [x.strip(" `").upper() for x in PK_RE.search(d).group(1).split(",")]
            elif u.startswith("FOREIGN KEY") or u.startswith("CONSTRAINT"):
                f = FK_RE.search(d)
                if f:
                    t["fks"].append(([x.strip(" `").upper() for x in f.group(1).split(",")], f.group(2).upper(),
                                     [x.strip(" `").upper() for x in f.group(3).split(",")]))
            elif u.startswith(("KEY ", "INDEX ", "UNIQUE", "CHECK")):
                continue
            else:
                c = COL_RE.match(d)
                if not c:
                    rep.err(name, "-", "D1", f"컬럼 정의를 읽지 못했습니다: '{d[:40]}'")
                    continue
                rest = c.group(5).upper()
                typ = c.group(2).upper()
                t["cols"][c.group(1).upper()] = {
                    "type": TYPE_ALIAS.get(typ, typ),
                    "len": int(c.group(3)) if c.group(3) else None,
                    "scale": int(c.group(4)) if c.group(4) else None,
                    "null": "NOT NULL" not in rest,
                }
                if "PRIMARY KEY" in rest:
                    t["pk"] = [c.group(1).upper()]
        tables[name] = t
    return tables


def load_dictionary():
    f = ROOT / "docs" / "표준사전.xlsx"
    if not f.exists():
        return None, None
    from openpyxl import load_workbook
    wb = load_workbook(f, read_only=True, data_only=True)

    def rows(sheet):
        r = list(wb[sheet].iter_rows(values_only=True))
        head = [str(h).strip() if h else "" for h in r[0]]
        return [dict(zip(head, x)) for x in r[1:] if x and x[0]]
    doms = {}
    for d in rows("도메인"):
        doms[str(d["도메인명"]).strip()] = (str(d["데이터타입"]).strip().upper(),
                                        int(d["길이"]) if d["길이"] not in (None, "") else None,
                                        int(d["소수점"]) if d["소수점"] not in (None, "") else None)
    terms = {str(t["물리명"]).strip(): str(t["도메인"]).strip() for t in rows("용어")}
    return terms, doms


def check(sql_text: str, rep: Report):
    sql = strip_comments(sql_text)
    for m in re.finditer(r"\bDROP\s+(DATABASE|SCHEMA|TABLE)\b", sql, re.I):
        rep.err("파일", "-", "D0", f"'{m.group(0)}' — ddl.sql 에는 지우는 명령을 넣지 않습니다 (실습 DB 초기화는 sql\\day05_reset.sql)")
    if not re.search(r"\bUSE\s+dstd_lab\b", sql, re.I):
        rep.warn("파일", "-", "D0", "USE dstd_lab; 이 없습니다 — 다른 DB 에 테이블이 생길 수 있습니다")
    tables = parse(sql, rep)
    if not tables:
        rep.err("파일", "-", "D1", "CREATE TABLE 문을 찾지 못했습니다")
        return {}
    for name, pk in TABLES.items():
        if name not in tables:
            rep.err(name, "-", "D2", "테이블이 없습니다")
            continue
        got = tables[name]["pk"]
        if not got:
            rep.err(name, "-", "D3", "PRIMARY KEY 가 없습니다")
        elif sorted(got) != sorted(pk):
            rep.err(name, "-", "D3", f"PK 가 ({', '.join(got)}) — 기대: ({', '.join(pk)})")
    for child, col, parent in FKS:
        if child not in tables:
            continue
        if not any(col in cols and ref == parent for cols, ref, _ in tables[child]["fks"]):
            rep.err(child, "-", "D4", f"FOREIGN KEY ({col}) REFERENCES {parent} 가 없습니다")
    # 만드는 순서 — 부모가 먼저
    order = [m.group(1).upper() for m in CREATE_RE.finditer(sql)]
    for child, col, parent in FKS:
        if child in order and parent in order and order.index(parent) > order.index(child):
            rep.err(child, "-", "D4", f"{parent} 보다 먼저 만들었습니다 — 부모 테이블이 없으면 FK 를 만들 수 없습니다")
    terms, doms = load_dictionary()
    for name, t in tables.items():
        for col, c in t["cols"].items():
            if terms is not None:
                if col not in terms:
                    rep.err(name, "-", "D5", f"컬럼 '{col}' — 표준 사전 「용어」 에 없는 물리명")
                elif terms[col] in doms:
                    typ, ln, sc = doms[terms[col]]
                    got = (c["type"], c["len"], c["scale"] if c["type"] == "DECIMAL" else None)
                    want = (typ, ln if typ != "INT" and typ != "DATE" else None, sc if typ == "DECIMAL" else None)
                    if got != want:
                        g = f"{got[0]}" + (f"({got[1]}" + (f",{got[2]}" if got[2] is not None else "") + ")" if got[1] else "")
                        w = f"{want[0]}" + (f"({want[1]}" + (f",{want[2]}" if want[2] is not None else "") + ")" if want[1] else "")
                        rep.err(name, "-", "D6", f"{col} 이 {g} — 도메인 '{terms[col]}' 은 {w}")
            if (name, col) in MUST_NULL and not c["null"]:
                rep.err(name, "-", "D7", f"{col} 이 NOT NULL — {MUST_NULL[(name, col)]} 이 적재에 실패합니다")
    return {"테이블": len(tables), "컬럼": sum(len(t["cols"]) for t in tables.values()),
            "FK": sum(len(t["fks"]) for t in tables.values())}


def run(path, quiet=False):
    path = Path(path)
    rep = Report()
    stats = check(path.read_text(encoding="utf-8"), rep)
    if not quiet:
        print(f"DDL 검사 — {path.name}")
        if stats:
            print("  " + " / ".join(f"{k} {v}" for k, v in stats.items()))
        print()
        for sev, where, _, rule, msg in rep.items:
            w = sum(2 if ord(ch) > 0x1100 else 1 for ch in where)
            print(f"[{sev}] {rule:<3} {where}{' ' * max(1, 10 - w)}{msg}")
        if rep.items:
            print()
        print("─" * 60)
        print(f"오류 {rep.n_err}건 · 경고 {rep.n_warn}건"
              + ("   → 통과" if rep.n_err == 0 else "   → 고칠 것이 있습니다"))
    return rep, stats


def main(argv=None):
    ap = argparse.ArgumentParser(description="DDL 정적 검사기")
    ap.add_argument("ddl", nargs="?", default=str(ROOT / "docs" / "ddl.sql"))
    a = ap.parse_args(argv)
    p = Path(a.ddl)
    if not p.exists():
        print(f"파일이 없습니다: {p}")
        print("templates\\ddl_템플릿.sql 을 docs\\ddl.sql 로 복사해 작성하십시오.")
        return 2
    rep, _ = run(p)
    return 1 if rep.n_err else 0


if __name__ == "__main__":
    sys.exit(main())

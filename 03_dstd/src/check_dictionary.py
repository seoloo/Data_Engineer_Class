# -*- coding: utf-8 -*-
"""
check_dictionary.py — 표준 사전·매핑표 규칙 검사기

표준은 사람끼리의 약속이지만, 약속이 지켜졌는지는 기계로 검사할 수 있다.
이 스크립트는 Google Sheets 에서 내려받은 xlsx 를 읽어 규칙 위반을 찾아낸다.

사용 예 (C:\\dstd 에서)
  .\\.venv\\Scripts\\python.exe src\\check_dictionary.py
  .\\.venv\\Scripts\\python.exe src\\check_dictionary.py templates\\표준사전_연습.xlsx
  .\\.venv\\Scripts\\python.exe src\\check_dictionary.py docs\\표준사전.xlsx --mapping docs\\매핑표.xlsx

기본 경로: docs\\표준사전.xlsx
행 번호는 스프레드시트에 보이는 행 번호(머리행 = 1)와 같다.
"""
from __future__ import annotations

import argparse
import csv
import json
import re
import sys
from pathlib import Path

from openpyxl import load_workbook

SHEETS = {
    "단어": ["단어명", "영문약어", "영문명", "분류어여부", "정의"],
    "도메인": ["도메인명", "분류어", "데이터타입", "길이", "소수점", "형식/허용값", "설명"],
    "용어": ["용어명", "물리명", "구성단어", "도메인", "코드그룹", "정의"],
    "코드": ["코드그룹", "코드그룹명", "코드값", "코드값명", "정렬순서", "설명"],
}
MAP_SHEETS = {
    "컬럼매핑": ["원천파일", "원천컬럼", "예시값", "표준용어", "물리명", "변환규칙", "비고"],
    "코드매핑": ["원천파일", "원천컬럼", "원천값", "표준코드그룹", "표준코드값", "처리"],
}
ABBR_RE = re.compile(r"^[A-Z][A-Z0-9]*$")


def find_root(start: Path) -> Path:
    for p in [start, *start.parents]:
        if (p / "src" / "make_dataset.py").exists():
            return p
    return start


ROOT = find_root(Path(__file__).resolve().parent)


# ─────────────────────────────────────────────────────────────
class Report:
    def __init__(self):
        self.items = []          # (심각도, 시트, 행, 규칙, 메시지)

    def err(self, sheet, row, rule, msg):
        self.items.append(("오류", sheet, row, rule, msg))

    def warn(self, sheet, row, rule, msg):
        self.items.append(("경고", sheet, row, rule, msg))

    @property
    def n_err(self):
        return sum(1 for i in self.items if i[0] == "오류")

    @property
    def n_warn(self):
        return sum(1 for i in self.items if i[0] == "경고")


def txt(v) -> str:
    """셀 값을 비교용 문자열로. 숫자 7.0 은 '7' 로."""
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip()


def read_sheet(wb, name, header, rep, required=True):
    """시트를 [{'_row': n, '_raw': {열: 원래셀값}, 열: 문자열}] 로 읽는다."""
    if name not in wb.sheetnames:
        if required:
            rep.err(name, "-", "S1", f"'{name}' 시트가 없습니다. 시트 이름을 확인하십시오")
        return []
    ws = wb[name]
    got = [txt(c.value) for c in ws[1]][: len(header)]
    if got != header:
        rep.err(name, 1, "S2", f"머리행이 다릅니다. 기대: {header} / 실제: {got}")
        return []
    rows = []
    for r in ws.iter_rows(min_row=2):
        vals = [c.value for c in r[: len(header)]]
        if all(v is None or str(v).strip() == "" for v in vals):
            continue
        d = {"_row": r[0].row, "_raw": dict(zip(header, vals))}
        d.update({h: txt(v) for h, v in zip(header, vals)})
        rows.append(d)
    return rows


# ─────────────────────────────────────────────────────────────
def check_dictionary(path: Path, rep: Report):
    wb = load_workbook(path, data_only=True)
    words = read_sheet(wb, "단어", SHEETS["단어"], rep)
    domains = read_sheet(wb, "도메인", SHEETS["도메인"], rep)
    terms = read_sheet(wb, "용어", SHEETS["용어"], rep)
    codes = read_sheet(wb, "코드", SHEETS["코드"], rep)

    # ── 단어 ───────────────────────────────────────────────
    by_name, by_abbr = {}, {}
    for w in words:
        n, a, r = w["단어명"], w["영문약어"], w["_row"]
        if n in by_name:
            rep.err("단어", r, "W1", f"단어 '{n}' — {by_name[n]['_row']}행과 중복")
        else:
            by_name[n] = w
        if not ABBR_RE.match(a):
            rep.err("단어", r, "W3", f"약어 '{a}' — 영문 대문자로 시작하고 대문자·숫자만 써야 함")
        if a in by_abbr:
            o = by_abbr[a]
            rep.err("단어", r, "W2", f"약어 '{a}' — '{o['단어명']}'({o['_row']}행)와 겹침")
        else:
            by_abbr[a] = w
        if w["분류어여부"] not in ("Y", "N"):
            rep.err("단어", r, "W4", f"분류어여부 '{w['분류어여부']}' — Y 또는 N 이어야 함")
    classwords = {n for n, w in by_name.items() if w["분류어여부"] == "Y"}

    # ── 도메인 ─────────────────────────────────────────────
    dom = {}
    for d in domains:
        n, r = d["도메인명"], d["_row"]
        if n in dom:
            rep.err("도메인", r, "D2", f"도메인 '{n}' — {dom[n]['_row']}행과 중복")
        dom[n] = d
        if d["분류어"] not in classwords:
            rep.err("도메인", r, "D1", f"분류어 '{d['분류어']}' — 단어 사전의 분류어(Y)에 없음")

    # ── 코드 ───────────────────────────────────────────────
    code_groups, seen_code = {}, {}
    for c in codes:
        g, v, r = c["코드그룹"], c["코드값"], c["_row"]
        code_groups.setdefault(g, set()).add(v)
        if (g, v) in seen_code:
            rep.err("코드", r, "C1", f"코드 {g}.{v} — {seen_code[(g, v)]}행과 중복")
        seen_code[(g, v)] = r
        if isinstance(c["_raw"]["코드값"], (int, float)):
            rep.warn("코드", r, "C2", f"코드값 '{v}' — 숫자 셀로 저장됨. "
                     "Google Sheets 가 '01' 을 1 로 바꿨을 수 있습니다. 열 서식을 '일반 텍스트'로")

    # ── 용어 ───────────────────────────────────────────────
    seen_t, seen_p = {}, {}
    for t in terms:
        n, p, r = t["용어명"], t["물리명"], t["_row"]
        parts = [x.strip() for x in t["구성단어"].split("+") if x.strip()]

        if n in seen_t:
            rep.err("용어", r, "T8", f"용어 '{n}' — {seen_t[n]}행과 중복")
        seen_t[n] = r
        if p in seen_p:
            rep.err("용어", r, "T8", f"물리명 '{p}' — {seen_p[p]}행과 중복")
        seen_p[p] = r

        if "".join(parts) != n:
            rep.err("용어", r, "T4", f"용어명 '{n}' — 구성단어 '{t['구성단어']}'를 이은 것과 다름")

        missing = [x for x in parts if x not in by_name]
        if missing:
            rep.err("용어", r, "T1", f"용어 '{n}' — 단어 {missing} 이(가) 단어 사전에 없음")
            continue                                   # 조합 불가 → 뒤 규칙 생략

        last = parts[-1] if parts else ""
        if last not in classwords:
            rep.err("용어", r, "T2", f"용어 '{n}' — 마지막 단어 '{last}' 는 분류어가 아님. "
                    "분류어(번호·일자·금액·코드·명·수량·여부 …)로 끝나야 함")
            continue

        expect = "_".join(by_name[x]["영문약어"] for x in parts)
        if p != expect:
            rep.err("용어", r, "T3", f"물리명 '{p}' — 단어 약어를 조합하면 '{expect}'")

        dname = t["도메인"]
        if dname not in dom:
            rep.err("용어", r, "T5", f"용어 '{n}' — 도메인 '{dname}' 이(가) 도메인 사전에 없음")
        elif dom[dname]["분류어"] != last:
            rep.err("용어", r, "T6", f"용어 '{n}' — '{last}'(으)로 끝나는데 도메인 '{dname}' 은(는) '{dom[dname]['분류어']}' 도메인")

        g = t["코드그룹"]
        if last == "코드":
            if not g:
                rep.err("용어", r, "T7", f"용어 '{n}' — 코드 용어인데 코드그룹이 비어 있음")
            elif g not in code_groups:
                rep.err("용어", r, "T7", f"코드그룹 '{g}' — 코드 사전에 없음")
        elif g:
            rep.warn("용어", r, "T7", f"용어 '{n}' — 코드 용어가 아닌데 코드그룹 '{g}' 이(가) 적혀 있음")

    stats = dict(단어=len(words), 도메인=len(domains), 용어=len(terms), 코드=len(codes))
    ctx = dict(terms={t["물리명"]: t for t in terms}, code_groups=code_groups)
    return stats, ctx


# ─────────────────────────────────────────────────────────────
def source_columns(data_dir: Path):
    """원천 3종의 실제 컬럼명을 읽는다."""
    cols = {}
    f = data_dir / "cust_master.csv"
    if f.exists():
        with f.open(encoding="euc-kr", newline="") as fh:
            cols[f.name] = next(csv.reader(fh))
    f = data_dir / "order_2026.csv"
    if f.exists():
        with f.open(encoding="utf-8-sig", newline="") as fh:
            cols[f.name] = next(csv.reader(fh))
    f = data_dir / "product_api_response.json"
    if f.exists():
        items = json.loads(f.read_text(encoding="utf-8"))["items"]
        cols[f.name] = list(items[0].keys()) if items else []
    return cols


def source_values(data_dir: Path, fname: str, col: str):
    """원천 파일 한 컬럼의 서로 다른 값(빈 값 제외)."""
    f = data_dir / fname
    if not f.exists():
        return None
    if fname.endswith(".json"):
        items = json.loads(f.read_text(encoding="utf-8"), parse_int=str, parse_float=str)["items"]
        vals = {str(it.get(col, "")).strip() for it in items}
    else:
        enc = "euc-kr" if fname == "cust_master.csv" else "utf-8-sig"
        with f.open(encoding=enc, newline="") as fh:
            vals = {(r.get(col) or "").strip() for r in csv.DictReader(fh)}
    vals.discard("")
    return vals


def check_mapping(path: Path, ctx, rep: Report, data_dir: Path | None):
    wb = load_workbook(path, data_only=True)
    colmap = read_sheet(wb, "컬럼매핑", MAP_SHEETS["컬럼매핑"], rep)
    codemap = read_sheet(wb, "코드매핑", MAP_SHEETS["코드매핑"], rep)
    terms, groups = ctx["terms"], ctx["code_groups"]

    listed = set()
    for m in colmap:
        r, p = m["_row"], m["물리명"]
        listed.add((m["원천파일"], m["원천컬럼"]))
        if p not in terms:
            rep.err("컬럼매핑", r, "M1", f"물리명 '{p}' — 용어 사전에 없음")
        elif terms[p]["용어명"] != m["표준용어"]:
            rep.err("컬럼매핑", r, "M2", f"물리명 '{p}' — 용어 사전에서는 '{terms[p]['용어명']}', 여기에는 '{m['표준용어']}'")

    if data_dir and data_dir.exists():
        actual = source_columns(data_dir)
        for fname, cols in actual.items():
            for c in cols:
                if (fname, c) not in listed:
                    rep.err("컬럼매핑", "-", "M3", f"{fname} 컬럼 '{c}' — 매핑표에 빠져 있음")
        for (fname, c) in sorted(listed):
            if fname in actual and c not in actual[fname]:
                rep.err("컬럼매핑", "-", "M3", f"{fname} — '{c}' 컬럼은 원천에 없음 (오타?)")

    seen = {}
    for m in codemap:
        r, g, v, how = m["_row"], m["표준코드그룹"], m["표준코드값"], m["처리"]
        key = (m["원천파일"], m["원천컬럼"], m["원천값"])
        if key in seen:
            rep.err("코드매핑", r, "M6", f"원천값 '{m['원천값']}' — {seen[key]}행과 중복. "
                    "Google Sheets 가 '01' 을 1 로 바꿔 '1' 과 겹쳤을 수 있습니다")
        seen[key] = r
        if isinstance(m["_raw"]["원천값"], (int, float)):
            rep.warn("코드매핑", r, "M5", f"원천값 '{m['원천값']}' — 숫자 셀로 저장됨. "
                     "선행 0 이 사라졌을 수 있습니다. 열 서식을 '일반 텍스트'로")
        if how == "매핑":
            if g not in groups:
                rep.err("코드매핑", r, "M4", f"코드그룹 '{g}' — 코드 사전에 없음")
            elif v not in groups[g]:
                rep.err("코드매핑", r, "M4", f"표준코드값 '{v}' — 코드 사전 {g}에 없음")
        elif how == "확인필요":
            if v:
                rep.err("코드매핑", r, "M4", f"'확인필요' 인데 표준코드값 '{v}' 이(가) 채워져 있음. "
                        "추측 매핑은 하지 않습니다")
        else:
            rep.err("코드매핑", r, "M4", f"처리 '{how}' — '매핑' 또는 '확인필요' 여야 함")

    # 원천에 실제로 나오는 코드값이 전부 매핑표에 있는가
    if data_dir and data_dir.exists():
        listed_vals = {}
        for m in codemap:
            listed_vals.setdefault((m["원천파일"], m["원천컬럼"]), set()).add(m["원천값"])
        for (fname, col), have in sorted(listed_vals.items()):
            actual = source_values(data_dir, fname, col)
            if actual is None:
                continue
            for v in sorted(actual - have):
                rep.err("코드매핑", "-", "M7", f"{fname} {col} — 원천값 '{v}' 매핑 누락 "
                        f"({data_dir.name} 기준)")

    return dict(컬럼매핑=len(colmap), 코드매핑=len(codemap))


# ─────────────────────────────────────────────────────────────
def print_report(rep: Report, title: str, stats: dict):
    print(title)
    print("  " + " / ".join(f"{k} {v}" for k, v in stats.items()))
    print()
    order = {"단어": 0, "도메인": 1, "용어": 2, "코드": 3, "컬럼매핑": 4, "코드매핑": 5}
    for sev, sheet, row, rule, msg in sorted(
            rep.items, key=lambda x: (order.get(x[1], 9), 0 if x[2] == "-" else x[2])):
        rtxt = f"{row}행" if row != "-" else "  - "
        print(f"[{sev}] {sheet:<4} {rtxt:>5}  {rule}  {msg}")
    if rep.items:
        print()
    print("─" * 60)
    print(f"오류 {rep.n_err}건 · 경고 {rep.n_warn}건"
          + ("   → 통과" if rep.n_err == 0 else "   → 고칠 것이 있습니다"))


def run(dict_path, mapping_path=None, data_dir=None, quiet=False):
    """노트북에서 부르기 위한 함수. (Report, stats) 를 돌려준다."""
    dict_path = Path(dict_path)
    rep = Report()
    stats, ctx = check_dictionary(dict_path, rep)
    if mapping_path:
        stats.update(check_mapping(Path(mapping_path), ctx, rep,
                                   Path(data_dir) if data_dir else ROOT / "data" / "sample"))
    if not quiet:
        title = f"표준 사전 검사 — {dict_path.name}"
        if mapping_path:
            title += f"  +  {Path(mapping_path).name}"
        print_report(rep, title, stats)
    return rep, stats


def main(argv=None):
    ap = argparse.ArgumentParser(description="표준 사전·매핑표 규칙 검사기")
    ap.add_argument("dictionary", nargs="?", default=str(ROOT / "docs" / "표준사전.xlsx"))
    ap.add_argument("--mapping", help="매핑표 xlsx 도 함께 검사")
    ap.add_argument("--data", help="원천 파일 폴더 (기본: data\\sample)")
    a = ap.parse_args(argv)
    p = Path(a.dictionary)
    if not p.exists():
        print(f"파일이 없습니다: {p}")
        print("Google Sheets 에서 '파일 → 다운로드 → Microsoft Excel(.xlsx)' 로 받아 이 위치에 저장하십시오.")
        return 2
    rep, _ = run(p, a.mapping, a.data)
    return 1 if rep.n_err else 0


if __name__ == "__main__":
    sys.exit(main())

# -*- coding: utf-8 -*-
"""
check_model.py — Mermaid ERD 검사기 (AS-IS 모델 · TO-BE 개념 모델)

그림은 사람이 보고, 규칙은 기계가 본다.
docs\\모델_ASIS.md, docs\\모델_TOBE.md 안의 ```mermaid erDiagram 블록을 읽어 검사한다.

실행 (C:\\dstd 에서)
  .\\.venv\\Scripts\\python.exe src\\check_model.py                      TO-BE (docs\\모델_TOBE.md)
  .\\.venv\\Scripts\\python.exe src\\check_model.py --asis               AS-IS (docs\\모델_ASIS.md)
  .\\.venv\\Scripts\\python.exe src\\check_model.py 파일.md [--asis]
"""
from __future__ import annotations

import argparse
import re
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
from check_dictionary import ROOT, Report, source_columns   # noqa: E402

CARD_L = r"\|o|\|\||\}o|\}\|"
CARD_R = r"o\||\|\||o\{|\|\{"
REL_RE = re.compile(rf"^(\S+)\s+({CARD_L})(--|\.\.)({CARD_R})\s+(\S+)\s*:\s*(.+)$")
ENT_RE = re.compile(r"^(\S+?)\s*\{\s*$")
ATTR_RE = re.compile(r'^(\S+)\s+(\S+)(?:\s+((?:PK|FK|UK)(?:\s*,\s*(?:PK|FK|UK))*))?(?:\s+"([^"]*)")?\s*$')
MIRROR = {"|o": "o|", "||": "||", "}o": "o{", "}|": "|{"}
CARD_TXT = {"|o": "0 또는 1", "||": "정확히 1", "}o": "0 이상", "}|": "1 이상"}

# ── TO-BE 기대값 ──────────────────────────────────────────────
TOBE_ENT = {"고객": ["CUST_NO"], "주문": ["ORD_NO"], "주문항목": ["ORD_NO", "ORD_SEQ"], "상품": ["PRD_NO"]}
# (부모, 자식, 부모 쪽 기대 카디널리티, 자식 쪽 기대 카디널리티, 이유)
TOBE_REL = [
    ("고객", "주문", "|o", "}o", "비회원 주문은 고객이 없습니다(업무설명서 2.1). 고객 쪽은 '0 또는 1'"),
    ("주문", "주문항목", "||", "}|", "항목 없는 주문은 없습니다. 항목 쪽은 '1 이상'"),
    ("상품", "주문항목", "||", "}o", "새로 등록한 상품은 아직 주문이 없습니다. 항목 쪽은 '0 이상'"),
]
ORDER_LEVEL = {"ORD_DT", "PAY_MTHD_CD", "DLV_FEE_AMT", "ORD_STAT_CD", "CNCL_DT"}
ITEM_LEVEL = {"PRD_NO", "ORD_QTY", "SALE_AMT"}
ASIS_ENT = {"cust_master": "cust_master.csv", "order_2026": "order_2026.csv",
            "product_api_response": "product_api_response.json"}


def parse(text: str, rep: Report):
    m = re.search(r"```mermaid\s*\n(.*?)```", text, re.S)
    body = m.group(1) if m else text
    lines = [ln.strip() for ln in body.splitlines()]
    if not any(ln.startswith("erDiagram") for ln in lines):
        rep.err("모델", "-", "E1", "erDiagram 블록을 찾지 못했습니다. ```mermaid 다음 줄에 erDiagram 이 있어야 합니다")
        return {}, []
    ents, rels, cur = {}, [], None
    for n, ln in enumerate(lines, 1):
        if not ln or ln.startswith("%%") or ln.startswith("erDiagram"):
            continue
        if cur is not None:
            if ln == "}":
                cur = None
                continue
            a = ATTR_RE.match(ln)
            if not a:
                rep.err(cur, "-", "E1", f"속성 줄을 읽지 못했습니다: '{ln}'  (형식: 타입 물리명 PK \"설명\")")
                continue
            keys = {k.strip() for k in (a.group(3) or "").split(",") if k.strip()}
            ents[cur].append({"type": a.group(1), "name": a.group(2), "keys": keys, "desc": a.group(4) or ""})
            continue
        e = ENT_RE.match(ln)
        if e:
            cur = e.group(1)
            ents.setdefault(cur, [])
            continue
        r = REL_RE.match(ln)
        if r:
            rels.append({"a": r.group(1), "ca": r.group(2), "line": r.group(3), "cb": r.group(4),
                         "b": r.group(5), "label": r.group(6).strip().strip('"')})
            continue
        rep.err("모델", "-", "E1", f"읽지 못한 줄: '{ln}'")
    for r in rels:
        for x in (r["a"], r["b"]):
            ents.setdefault(x, [])
    return ents, rels


def find_rel(rels, parent, child):
    """방향과 상관없이 parent–child 관계를 찾아 (부모 쪽 카디널리티, 자식 쪽 카디널리티, 선) 로 돌려준다."""
    inv = {v: k for k, v in MIRROR.items()}
    for r in rels:
        if r["a"] == parent and r["b"] == child:
            return r["ca"], inv[r["cb"]], r["line"]
        if r["a"] == child and r["b"] == parent:
            return inv[r["cb"]], r["ca"], r["line"]
    return None


def load_terms():
    f = ROOT / "docs" / "표준사전.xlsx"
    if not f.exists():
        return None
    from openpyxl import load_workbook
    ws = load_workbook(f, read_only=True, data_only=True)["용어"]
    rows = list(ws.iter_rows(values_only=True))
    ix = [str(h).strip() if h else "" for h in rows[0]].index("물리명")
    return {str(r[ix]).strip() for r in rows[1:] if r and r[ix]}


def check_tobe(ents, rels, rep: Report):
    for name, pk in TOBE_ENT.items():
        if name not in ents:
            rep.err(name, "-", "E2", "엔터티가 없습니다")
            continue
        got = [a["name"] for a in ents[name] if "PK" in a["keys"]]
        if not got:
            rep.err(name, "-", "E3", "식별자(PK)가 없습니다")
        elif sorted(got) != sorted(pk):
            rep.err(name, "-", "E4", f"식별자가 {' + '.join(got)} — 기대: {' + '.join(pk)}")
    terms = load_terms()
    for name, attrs in ents.items():
        for a in attrs:
            if terms is not None and a["name"] not in terms:
                rep.err(name, "-", "E5", f"속성 '{a['name']}' — 표준 사전 「용어」 에 없는 물리명")
        names = {a["name"] for a in attrs}
        if name == "주문항목":
            for x in sorted(names & ORDER_LEVEL):
                rep.err(name, "-", "E6", f"'{x}' 는 주문 단위 속성입니다 — 주문 엔터티로 (셀 4 판정)")
        if name == "주문":
            for x in sorted(names & ITEM_LEVEL):
                rep.err(name, "-", "E6", f"'{x}' 는 항목 단위 속성입니다 — 주문항목 엔터티로 (셀 4 판정)")
    for parent, child, want_p, want_c, why in TOBE_REL:
        if parent not in ents or child not in ents:
            continue
        got = find_rel(rels, parent, child)
        if not got:
            rep.err("관계", "-", "E7", f"{parent} — {child} 관계가 없습니다")
            continue
        cp, cc, _ = got
        if cp != want_p:
            rep.err("관계", "-", "E8", f"{parent}–{child}: {parent} 쪽이 '{CARD_TXT[cp]}' — {why}")
        if cc != want_c:
            rep.err("관계", "-", "E8", f"{parent}–{child}: {child} 쪽이 '{CARD_TXT[cc]}' — {why}")
        pk = TOBE_ENT[parent]
        fks = {a["name"] for a in ents[child] if "FK" in a["keys"]}
        for k in pk:
            if k not in fks:
                rep.err(child, "-", "E9", f"'{k}' 가 FK 로 없습니다 — {parent} 와 이어 주는 속성")
    # 같은 속성이 FK 가 아닌데 두 엔터티에 — 중복 저장
    seen = {}
    for name, attrs in ents.items():
        for a in attrs:
            if "FK" in a["keys"]:
                continue
            if a["name"] in seen and seen[a["name"]] != name:
                rep.warn(name, "-", "W1", f"'{a['name']}' 가 {seen[a['name']]} 에도 있습니다 — 같은 값을 두 곳에 저장하면 어긋납니다")
            seen.setdefault(a["name"], name)
    return {"엔터티": len(ents), "관계": len(rels), "속성": sum(len(v) for v in ents.values())}


def check_asis(ents, rels, rep: Report, data_dir: Path):
    cols = source_columns(data_dir)
    for ent, fname in ASIS_ENT.items():
        if ent not in ents:
            rep.err(ent, "-", "A1", f"엔터티가 없습니다 — {fname} 를 그대로 옮긴 엔터티 (이름: {ent})")
            continue
        have = {a["name"] for a in ents[ent]}
        for c in cols.get(fname, []):
            if c not in have:
                rep.err(ent, "-", "A2", f"원천 컬럼 '{c}' 가 빠졌습니다 — AS-IS 는 있는 그대로")
        for c in sorted(have - set(cols.get(fname, []))):
            rep.err(ent, "-", "A3", f"'{c}' 는 {fname} 에 없는 컬럼입니다 — AS-IS 에 표준 이름을 쓰지 마십시오")
    for r in rels:
        if r["line"] == "--":
            rep.warn("관계", "-", "A4", f"{r['a']}–{r['b']} 가 실선(--) — 원천에는 참조 검사가 없었습니다. 점선(..)으로")
    if not rels:
        rep.warn("관계", "-", "A5", "관계가 하나도 없습니다 — 원천끼리 값으로 이어지는 곳(점선)을 표시하십시오")
    return {"엔터티": len(ents), "관계": len(rels), "속성": sum(len(v) for v in ents.values())}


def run(path, asis=False, data_dir=None, quiet=False):
    path = Path(path)
    rep = Report()
    ents, rels = parse(path.read_text(encoding="utf-8"), rep)
    if ents:
        stats = (check_asis(ents, rels, rep, Path(data_dir) if data_dir else ROOT / "data" / "sample")
                 if asis else check_tobe(ents, rels, rep))
    else:
        stats = {}
    if not quiet:
        print(f"{'AS-IS' if asis else 'TO-BE'} 모델 검사 — {path.name}")
        if stats:
            print("  " + " / ".join(f"{k} {v}" for k, v in stats.items()))
        print()
        for sev, where, _, rule, msg in rep.items:
            w = sum(2 if ord(ch) > 0x1100 else 1 for ch in where)
            print(f"[{sev}] {rule:<3} {where}{' ' * max(1, 21 - w)}{msg}")
        if rep.items:
            print()
        print("─" * 60)
        print(f"오류 {rep.n_err}건 · 경고 {rep.n_warn}건"
              + ("   → 통과" if rep.n_err == 0 else "   → 고칠 것이 있습니다"))
    return rep, stats


def main(argv=None):
    ap = argparse.ArgumentParser(description="Mermaid ERD 검사기")
    ap.add_argument("model", nargs="?")
    ap.add_argument("--asis", action="store_true", help="AS-IS 모델로 검사")
    ap.add_argument("--data", help="원천 폴더 (AS-IS 검사용, 기본 data\\sample)")
    a = ap.parse_args(argv)
    p = Path(a.model) if a.model else ROOT / "docs" / ("모델_ASIS.md" if a.asis else "모델_TOBE.md")
    if not p.exists():
        print(f"파일이 없습니다: {p}")
        print("templates 폴더의 템플릿을 docs 로 복사해 작성하십시오.")
        return 2
    rep, _ = run(p, a.asis, a.data)
    return 1 if rep.n_err else 0


if __name__ == "__main__":
    sys.exit(main())
